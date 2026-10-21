"""
Sinh bảng theo dõi SV4 mới (docs/sv4/01_output/TRACKING_SV4.xlsx + TRACKING_SV4.md).

Sắp xếp theo đúng hành trình của tiền (yêu cầu của thầy): khách bắt đầu trả tiền -> SePay báo về ->
kế toán ghi nhận, xử lý ngoại lệ, chứng từ -> cuối kỳ đối soát, hoàn tiền, báo cáo.
Giữ nguyên mã UC047 - UC060 của SV4 (không lấn sang mã của thành viên khác), đổi thứ tự, tên và mô tả cho khớp code.
Cùng định dạng bảng chung project_tracking_group2.xlsx: header vàng, Arial 10, viền mảnh, ma trận CRUD.

Chạy:  python docs/sv4/03_tools/gen_tracking.py
"""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path(__file__).resolve().parents[1] / "01_output"

# ---------------------------------------------------------------------------
# 1. USE CASE (thứ tự theo luồng)
# ---------------------------------------------------------------------------
# (bước, mã, tên, actor, mô tả, iter, màn hình, route)
USE_CASES = [
    (1, "UC047", "Create Payment Order", "Reader, Advertiser",
     "Module nguồn (SV1 booking, SV3 mua bài / gói) tạo đơn PENDING mã LP..., giữ 15 phút, rồi chuyển khách sang trang thanh toán chung. "
     "Mọi đơn mua đều trả bằng Ví LocalPress: ví đủ thì 1 nút trừ ví; ví thiếu thì trang tự tạo QR nạp bù đúng phần còn thiếu (đơn TOPUP gắn mã đơn), "
     "tiền nạp về thì hệ thống tự trừ ví trả đơn. Nút Nạp tiền trên header tạo đơn TOPUP tự do.",
     "Iter 1", "II.4.1, II.4.2", "/payment/:orderCode, /wallet, /wallet/topup"),
    (2, "UC050", "Check Payment Status", "Reader, Advertiser, Payment Gateway",
     "Trang thanh toán hỏi trạng thái 3 giây/lần; nút \"Tôi đã chuyển khoản, kiểm tra lại\" hỏi thẳng API SePay cho đơn đó. "
     "Không bao giờ cấp quyền dựa vào trình duyệt (Rule 4).",
     "Iter 1", "II.4.1", "/payment/:orderCode"),
    (3, "UC048", "Verify SePay Payment", "Payment Gateway (SePay)",
     "Xác thực webhook bằng API Key (sai khóa trả 401), đọc payload SePay (id, transferAmount, content, referenceCode), "
     "tìm mã đơn trong nội dung chuyển khoản, so số tiền với đơn. Mỗi lần nhận ghi 1 dòng nhật ký kèm payload gốc.",
     "Iter 1", "II.4.3", "/backoffice/finance/sepay"),
    (4, "UC049", "Handle Payment Webhook (Idempotent)", "Payment Gateway (SePay)",
     "Xử lý tiền về không trùng: sepay_id UNIQUE, đơn đã SUCCESS trả ALREADY_PROCESSED. Tiền nạp -> cộng ví, nạp bù thì tự trả đơn gắn kèm "
     "(đủ -> SUCCESS + PaymentSuccessEvent; chưa đủ -> giữ trong ví). Chuyển thẳng vào mã đơn (cách cũ): đủ/thừa -> SUCCESS, thiếu hoặc về muộn -> vào ví. Luôn trả {success:true} khi đã xử lý để SePay ngừng gửi lại. "
     "Có đồng bộ dự phòng qua API SePay khi không nhận được webhook.",
     "Iter 1", "II.4.3", "POST /api/v1/finance/webhook/sepay"),
    (5, "UC051", "Handle Unrecorded Payment Complaint", "Reader, Advertiser, Accountant",
     "Khách đã chuyển tiền mà sau 15 phút chưa được ghi nhận thì gửi khiếu nại (mã GD ngân hàng UNIQUE). Kế toán đối chiếu sao kê rồi "
     "Ghi nhận vào đơn (chạy lại đúng luồng webhook), Cộng vào ví, hoặc Từ chối có lý do. Thay cho duyệt chuyển khoản thủ công.",
     "Iter 2", "II.4.4, II.4.5", "/wallet/complaints, /backoffice/finance/complaints"),
    (6, "UC057", "Generate Financial Documents", "Accountant",
     "Kế toán tra danh sách đơn, xem chi tiết các lần thanh toán và in hóa đơn điện tử / chứng từ cho đơn đã thanh toán.",
     "Iter 2", "II.4.6", "/backoffice/finance/orders, /backoffice/finance/orders/:id"),
    (7, "UC058", "View Financial Ledger", "Accountant",
     "Sổ quỹ: bút toán thu (Có) và hoàn (Nợ), số dư ròng tính từ các dòng. Tiền nạp ví là công nợ phải trả khách, không tính doanh thu (Rule 16).",
     "Iter 2", "II.4.7", "/backoffice/finance/ledger"),
    (8, "UC052", "View Cash Flow Dashboard", "Accountant, Chief Accountant",
     "Doanh thu thực thu, biểu đồ theo ngày, việc cần xử lý (hoàn tiền chờ duyệt, đơn chờ, chênh lệch đối soát, khiếu nại).",
     "Iter 3", "II.4.8", "/backoffice/finance"),
    (9, "UC053", "Manage Advertiser Receivables", "Accountant, Advertiser",
     "Công nợ quảng cáo theo chiến dịch: PAID_IN_FULL / IN_TERM / OVERDUE, số ngày quá hạn, nhắc nợ.",
     "Iter 3", "II.4.9", "(API GET /finance/receivables, chưa có màn hình)"),
    (10, "UC054", "Reconcile Payment Mismatch", "Accountant, Chief Accountant, Bank Statement",
     "Đối soát kỳ: tổng hệ thống vs sao kê, xử lý từng khoản lệch có ghi chú, kế toán trưởng chốt kỳ.",
     "Iter 3", "II.4.10", "/backoffice/finance/reconciliation"),
    (11, "UC055", "Process Refund Request (Four-Eyes)", "Accountant, Chief Accountant",
     "Mắt 1 kế toán viên lập đề xuất (tối đa số đã thu, tính cả đề xuất đang chờ - Rule 11), mắt 2 kế toán trưởng duyệt; người lập không tự duyệt.",
     "Iter 3", "II.4.11", "/backoffice/finance/refunds"),
    (12, "UC056", "Adjust Reading or Campaign Rights", "System, SV1 / SV3 Modules",
     "Hoàn tất hoàn tiền -> ghi giao dịch REFUND (đơn trả bằng ví thì hoàn vào ví) và phát RefundCompletedEvent; "
     "SV1/SV3 tự thu hồi quyền khi hoàn toàn bộ (Rule 17).",
     "Iter 3", "II.4.11", "(sự kiện RefundCompletedEvent)"),
    (13, "UC059", "Export Financial Report", "Chief Accountant",
     "Xuất CSV giao dịch (UTF-8 có BOM cho Excel) theo khoảng ngày.",
     "Iter 3", "II.4.12", "(API GET /finance/reports/export, chưa có màn hình)"),
    (14, "UC060", "View Financial Alerts", "Accountant",
     "Cảnh báo: hoàn tiền chờ duyệt, khiếu nại mở, công nợ quá hạn, webhook lỗi.",
     "Iter 3", "II.4.12", "(API GET /finance/alerts, chưa có màn hình)"),
]

# ---------------------------------------------------------------------------
# 2. MÀN HÌNH / CHỨC NĂNG (ma trận CRUD)
# ---------------------------------------------------------------------------
ENTITIES = [
    "users", "categories", "category_follows", "articles", "article_versions", "comments", "saved_articles",
    "subscription_plans", "subscriptions", "article_purchases", "advertisers", "ad_slots", "ad_campaigns",
    "ad_creatives", "ad_stats", "transactions", "refund_requests", "notifications", "audit_logs",
    # Bảng mở rộng của SV4
    "wallets", "wallet_entries", "payment_complaints", "sepay_webhook_logs",
]

# (tên, loại, {bảng: CRUD}, UC, iter, SRS, status, ghi chú, route)
FUNCTIONS = [
    ("Payment Checkout", "Screen",
     {"transactions": "CRU", "wallets": "RU", "wallet_entries": "C"},
     "UC047, UC050", "Iter 1", "II.4.1", "Done",
     "[Đầu quy trình] Trang thanh toán chung, mọi đơn trả bằng ví: ví đủ -> 1 nút trừ ví (khóa dòng đơn + ví, 402 thiếu tiền, 410 hết hạn); "
     "ví thiếu -> QR nạp bù đúng phần thiếu (shortfall-topup, dùng lại mã còn hạn), tiền về tự trả đơn. Polling 3 giây, kiểm tra lại qua API SePay. "
     "Không phải chủ đơn -> 403.",
     "/payment/:orderCode"),
    ("LocalPress Wallet and Top-up", "Screen",
     {"transactions": "C", "wallets": "R", "wallet_entries": "R"},
     "UC047, UC050", "Iter 1", "II.4.2", "Done",
     "[Đầu quy trình] Số dư + lịch sử ví; nạp tiền 10.000 - 50.000.000 ₫ (bội số 1.000) tạo đơn TOPUP rồi sang trang thanh toán chung. "
     "Tiền nạp là công nợ phải trả khách, không phải doanh thu.",
     "/wallet, /wallet/topup"),
    ("SePay Webhook and Payment Log", "Screen + API",
     {"transactions": "RU", "wallets": "U", "wallet_entries": "C", "sepay_webhook_logs": "CR"},
     "UC048, UC049", "Iter 1", "II.4.3", "Done",
     "[Đầu quy trình] Webhook SePay (API Key, idempotent theo sepay_id) + nhật ký từng lần tiền về kèm payload gốc; "
     "đồng bộ dự phòng qua API SePay. Kế toán xem nguồn Webhook / Đồng bộ API / Giả lập. Cần tunnel để SePay gọi vào máy local.",
     "/backoffice/finance/sepay"),
    ("Payment Complaint (Customer)", "Screen",
     {"transactions": "R", "payment_complaints": "CR"},
     "UC051", "Iter 2", "II.4.4", "Done",
     "[Giữa quy trình] Khách gửi khiếu nại đã chuyển tiền chưa được ghi nhận: số tiền, thời gian, mã GD ngân hàng (UNIQUE, 409 nếu trùng), "
     "mã đơn (404 nếu không có, 403 nếu của người khác); theo dõi phản hồi của kế toán.",
     "/wallet/complaints"),
    ("Payment Complaint Queue", "Screen",
     {"transactions": "RU", "wallets": "U", "wallet_entries": "C", "payment_complaints": "RU"},
     "UC051", "Iter 2", "II.4.5", "Done",
     "[Giữa quy trình] Kế toán đối chiếu sao kê: Ghi nhận vào đơn (chạy lại hàm webhook), Cộng vào ví (TOPUP CMP-...), Từ chối bắt buộc lý do. "
     "Khóa dòng khiếu nại, xử lý lại -> 409. Không có nút đánh dấu PAID.",
     "/backoffice/finance/complaints"),
    ("Payment Orders and Financial Documents", "Screen",
     {"users": "R", "transactions": "R"},
     "UC057", "Iter 2", "II.4.6", "Done",
     "[Giữa quy trình] Danh sách đơn (lọc loại, trạng thái, tìm kiếm), chi tiết các lần thanh toán, hóa đơn điện tử in / lưu PDF cho đơn đã thanh toán, "
     "nút lập đề xuất hoàn tiền.",
     "/backoffice/finance/orders, /backoffice/finance/orders/:id"),
    ("Financial Ledger", "Screen",
     {"transactions": "R"},
     "UC058", "Iter 2", "II.4.7", "Done",
     "[Giữa quy trình] Sổ quỹ: thu (Có), hoàn (Nợ), số dư ròng; loại TOPUP khỏi doanh thu (Rule 16).",
     "/backoffice/finance/ledger"),
    ("Cash Flow Dashboard", "Screen",
     {"transactions": "R", "refund_requests": "R"},
     "UC052", "Iter 3", "II.4.8", "Done",
     "[Cuối quy trình] Doanh thu thực thu, biểu đồ 14 ngày, hàng việc cần xử lý.",
     "/backoffice/finance"),
    ("Advertiser Receivables", "Screen",
     {"advertisers": "R", "ad_campaigns": "R", "transactions": "R"},
     "UC053", "Iter 3", "II.4.9", "In Progress",
     "[Cuối quy trình] Công nợ quảng cáo theo chiến dịch. Đã có API GET /finance/receivables, chưa có màn hình, chưa có nhắc nợ.",
     "(chưa có)"),
    ("Payment Reconciliation", "Screen",
     {"transactions": "RU"},
     "UC054", "Iter 3", "II.4.10", "In Progress",
     "[Cuối quy trình] Đối soát kỳ, xử lý khoản lệch, chốt kỳ. Sao kê đang giả lập bằng tổng hệ thống, khoản lệch lưu trong bộ nhớ.",
     "/backoffice/finance/reconciliation"),
    ("Refund and Rights Management", "Screen",
     {"transactions": "CRU", "refund_requests": "CRU", "wallets": "U", "wallet_entries": "C"},
     "UC055, UC056", "Iter 3", "II.4.11", "Done",
     "[Cuối quy trình] Hoàn tiền 4 mắt, Refund Cap, giao dịch REFUND, hoàn vào ví cho đơn trả bằng ví, phát RefundCompletedEvent. "
     "Listener thu hồi quyền của SV1 / SV3 chưa có.",
     "/backoffice/finance/refunds"),
    ("Financial Reports and Alerts", "Screen",
     {"transactions": "R", "refund_requests": "R", "payment_complaints": "R", "sepay_webhook_logs": "R"},
     "UC059, UC060", "Iter 3", "II.4.12", "In Progress",
     "[Cuối quy trình] Xuất CSV, cảnh báo. Đã có API, chưa có màn hình, CSV chưa lọc theo ngày.",
     "(chưa có)"),
]

# ---------------------------------------------------------------------------
# 3. LUỒNG NGHIỆP VỤ ĐẦU -> CUỐI
# ---------------------------------------------------------------------------
# (bước, ai làm, việc, màn hình, UC, iter, swimlane)
FLOW = [
    (1, "Độc giả / Doanh nghiệp", "Bấm Mua bài, Mua gói, Thanh toán booking (màn hình của SV3 / SV1); module nguồn tạo bản ghi chờ rồi gọi POST /finance/orders",
     "SV3 / SV1 -> II.4.1", "UC047", "Iter 1", "01"),
    (2, "Hệ thống", "Tạo đơn PENDING mã LP..., hết hạn 15 phút, chuyển khách sang /payment/:orderCode", "II.4.1", "UC047", "Iter 1", "01"),
    (3, "Hệ thống", "Kiểm tra ví: đủ tiền thì hiện 1 nút Thanh toán bằng ví (trừ ví, SUCCESS, PaymentSuccessEvent)", "II.4.1", "UC047", "Iter 1", "01"),
    (4, "Khách", "Ví thiếu: trang tự tạo QR nạp bù đúng phần thiếu (đơn TOPUP gắn mã đơn); khách quét QR chuyển khoản. Nạp tự do: nút Nạp tiền trên header", "II.4.1, II.4.2", "UC047", "Iter 1", "01, 02"),
    (5, "SePay -> Hệ thống", "Ngân hàng báo có, SePay gọi webhook: xác thực API Key, ghi nhật ký, cộng ví, nạp bù thì tự trừ ví trả đơn; gửi lại không xử lý trùng", "II.4.3", "UC048, UC049", "Iter 1", "01"),
    (6, "Trang thanh toán", "Hỏi trạng thái 3 giây/lần; nút Kiểm tra lại hỏi thẳng API SePay; đơn SUCCESS -> module nguồn mở quyền qua PaymentSuccessEvent", "II.4.1", "UC050", "Iter 1", "01"),
    (7, "Khách", "Đã chuyển hơn 15 phút chưa được ghi nhận: gửi khiếu nại kèm mã GD ngân hàng", "II.4.4", "UC051", "Iter 2", "03"),
    (8, "Kế toán", "Đối chiếu sao kê: ghi nhận vào đơn, cộng vào ví, hoặc từ chối có lý do", "II.4.5", "UC051", "Iter 2", "03"),
    (9, "Kế toán", "Tra đơn, xem các lần thanh toán, in hóa đơn điện tử", "II.4.6", "UC057", "Iter 2", "06"),
    (10, "Kế toán", "Xem sổ quỹ thu / hoàn, số dư ròng", "II.4.7", "UC058", "Iter 2", "06"),
    (11, "Kế toán / Kế toán trưởng", "Theo dõi dòng tiền và việc cần xử lý trên dashboard", "II.4.8", "UC052", "Iter 3", "06"),
    (12, "Kế toán", "Theo dõi công nợ quảng cáo, nhắc nợ", "II.4.9", "UC053", "Iter 3", "04"),
    (13, "Kế toán -> Kế toán trưởng", "Đối soát kỳ với sao kê, xử lý khoản lệch, chốt kỳ", "II.4.10", "UC054", "Iter 3", "04"),
    (14, "Kế toán -> Kế toán trưởng", "Hoàn tiền 4 mắt; hoàn tất thì SV1/SV3 thu hồi quyền qua RefundCompletedEvent", "II.4.11", "UC055, UC056", "Iter 3", "05"),
    (15, "Kế toán trưởng", "Xuất báo cáo, xem cảnh báo", "II.4.12", "UC059, UC060", "Iter 3", "06"),
]

# ---------------------------------------------------------------------------

HEADER_FILL = PatternFill("solid", fgColor="FFFF00")
ITER_FILL = {"Iter 1": "E8F0FE", "Iter 2": "FFF4E5", "Iter 3": "F1F3F4"}
THIN = Side(style="thin", color="999999")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)


def write_sheet(ws, headers, rows, widths, iter_col=None, height=None):
    ws.append(headers)
    for c in ws[1]:
        c.font = BOLD
        c.fill = HEADER_FILL
        c.border = BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    ws.row_dimensions[1].height = 30
    for r in rows:
        ws.append(list(r))
    for row in ws.iter_rows(min_row=2):
        fill = None
        if iter_col is not None:
            fill = PatternFill("solid", fgColor=ITER_FILL.get(row[iter_col].value, "FFFFFF"))
        for c in row:
            c.font = FONT
            c.border = BORDER
            c.alignment = Alignment(wrap_text=True, vertical="top")
            if fill:
                c.fill = fill
        if height:
            ws.row_dimensions[row[0].row].height = height
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "B2"


def main():
    wb = Workbook()

    # --- Use Case ---
    ws = wb.active
    ws.title = "Use Case"
    rows = [(f"B{b}", uc, name, actor, "SV4", desc, it, scr, route, "Huy")
            for (b, uc, name, actor, desc, it, scr, route) in USE_CASES]
    write_sheet(ws, ["Bước", "ID", "UC", "Actor", "Flow", "Description", "Iter", "Màn hình (SRS)", "Route / API", "In Charge"],
                rows, [7, 9, 34, 26, 7, 80, 8, 14, 34, 10], iter_col=6)
    dv = DataValidation(type="list", formula1='"SV1,SV2,SV3,SV4,SV5"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"E2:E{len(rows) + 1}")

    # --- Functions ---
    ws = wb.create_sheet("Functions")
    headers = ["Screen/Function", "Type"] + ENTITIES + ["Use Cases", "Plan", "In Charge", "SRS", "SDS", "Status", "Route", "Notes / Updates"]
    frows = []
    for (name, typ, crud, ucs, it, srs, status, note, route) in FUNCTIONS:
        frows.append([f"{srs} {name}", typ] + [crud.get(e) for e in ENTITIES] + [ucs, it, "SV4 (Huy)", srs, srs, status, route, note])
    widths = [36, 11] + [10] * len(ENTITIES) + [16, 9, 11, 9, 9, 12, 30, 80]
    write_sheet(ws, headers, frows, widths, iter_col=len(ENTITIES) + 3, height=60)
    for col in range(3, 3 + len(ENTITIES)):
        for row in ws.iter_rows(min_row=2, min_col=col, max_col=col):
            row[0].alignment = Alignment(horizontal="center", vertical="top")
    st = DataValidation(type="list", formula1='"To Do,In Progress,Done"', allow_blank=True)
    ws.add_data_validation(st)
    status_col = get_column_letter(len(ENTITIES) + 8)
    st.add(f"{status_col}2:{status_col}{len(frows) + 1}")

    # --- Luồng nghiệp vụ ---
    ws = wb.create_sheet("Luồng nghiệp vụ")
    write_sheet(ws, ["Bước", "Ai làm", "Việc", "Màn hình", "UC", "Iter", "Swimlane"],
                FLOW, [7, 24, 90, 16, 14, 8, 10], iter_col=5)

    # --- Tổng hợp Iter ---
    ws = wb.create_sheet("Tổng hợp Iter")
    summary = []
    for it in ("Iter 1", "Iter 2", "Iter 3"):
        fs = [f for f in FUNCTIONS if f[4] == it]
        ucs = sorted({u.strip() for f in fs for u in f[3].split(",")})
        summary.append((it, len(fs), "\n".join(f"{f[5]} {f[0]}" for f in fs), ", ".join(ucs),
                        sum(1 for f in fs if f[6] == "Done")))
    write_sheet(ws, ["Iter", "Số màn hình", "Màn hình", "Use case", "Đã xong"], summary, [8, 12, 50, 40, 10], iter_col=0, height=80)

    out = OUT_DIR / "TRACKING_SV4.xlsx"
    wb.save(out)
    print("Đã ghi", out)

    # --- Bản markdown đi kèm ---
    md = ["# Tracking SV4 (Huy) — sắp theo hành trình của tiền", "",
          "> Sinh bởi `docs/sv4/03_tools/gen_tracking.py`, cùng nội dung với `TRACKING_SV4.xlsx`.",
          "> Thứ tự: khách bắt đầu trả tiền → SePay báo về → kế toán ghi nhận, xử lý ngoại lệ, chứng từ → cuối kỳ đối soát, hoàn tiền, báo cáo.",
          "> Mã UC giữ nguyên UC047 - UC060; số màn hình SRS/SDS đánh lại II.4.1 - II.4.12 theo thứ tự luồng.", ""]
    for it, n, scr, ucs, done in summary:
        md.append(f"- **{it}**: {n} màn hình ({done} xong). UC: {ucs}")
    md += ["", "## 1. Use case theo thứ tự luồng", "",
           "| Bước | ID | UC | Actor | Iter | Màn hình | Route / API |", "|---|---|---|---|---|---|---|"]
    for (b, uc, name, actor, desc, it, scr, route) in USE_CASES:
        md.append(f"| {b} | {uc} | **{name}**<br>{desc} | {actor} | {it} | {scr} | `{route}` |")
    md += ["", "## 2. Màn hình / chức năng", "",
           "| SRS | Màn hình | Iter | UC | Status | Route | Bảng (CRUD) | Ghi chú |", "|---|---|---|---|---|---|---|---|"]
    for (name, typ, crud, ucs, it, srs, status, note, route) in FUNCTIONS:
        crud_txt = ", ".join(f"{k} {v}" for k, v in crud.items())
        md.append(f"| {srs} | {name} | {it} | {ucs} | {status} | `{route}` | {crud_txt} | {note} |")
    md += ["", "## 3. Luồng nghiệp vụ đầu → cuối", "",
           "| Bước | Ai làm | Việc | Màn hình | UC | Iter | Swimlane |", "|---|---|---|---|---|---|---|"]
    for r in FLOW:
        md.append("| " + " | ".join(str(x) for x in r) + " |")
    md += ["", "## 4. So với tracking cũ (project_tracking_group2.xlsx)", "",
           "| Cũ | Mới | Lý do |", "|---|---|---|",
           "| Iter 1: Financial Documents and Ledger (II.4.1) | Iter 2: II.4.6, II.4.7 | Chứng từ và sổ quỹ chỉ có dữ liệu sau khi tiền đã về, nên ở giữa luồng |",
           "| Iter 1: Manual Bank Transfer Confirmation (II.4.2, UC051) | Iter 2: II.4.4 + II.4.5 Payment Complaint | Đã bỏ duyệt chuyển khoản thủ công; UC051 thành khiếu nại có đối chiếu sao kê |",
           "| Iter 2: Payment Processing (II.4.3) | Iter 1: II.4.1 Payment Checkout + II.4.3 SePay Webhook and Payment Log | Thanh toán là điểm bắt đầu của luồng tiền |",
           "| (không có) | Iter 1: II.4.2 LocalPress Wallet and Top-up | Ví thêm ở V3; từ 22/10/2026 mọi đơn trả qua ví, thiếu thì nạp bù (V6 `topup_for_order`) |",
           "| Iter 2: Advertiser Receivables (II.4.4) | Iter 3: II.4.9 | Công nợ là việc cuối kỳ, chưa có màn hình |",
           "| Iter 3: 4 màn hình | Iter 3: 5 màn hình (II.4.8 - II.4.12) | Giữ nguyên nội dung, đánh lại số |",
           "", "Cần cập nhật theo: số mục SRS/SDS của SV4 trong `Template1_SRS_Gr2.docx` và dòng SV4 trong `project_tracking_group2.xlsx` (nếu nhóm đồng ý dùng bản này).", ""]
    (OUT_DIR / "TRACKING_SV4.md").write_text("\n".join(md), encoding="utf-8")
    print("Đã ghi", OUT_DIR / "TRACKING_SV4.md")


if __name__ == "__main__":
    main()
